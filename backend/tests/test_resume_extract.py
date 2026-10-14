import io
import json
from unittest.mock import patch

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.core.config import Settings, get_settings
from app.main import app
from app.schemas.resume import ResumeData
from app.services.exporter.pdf_exporter import generate_resume_pdf
from tests.test_tailor import MOCK_TAILORED_PAYLOAD, MockGenAIClient, get_mock_settings

client = TestClient(app)


def create_sample_pdf(text: str = "Jane Doe\nSenior Backend Engineer\nOver 8 years experience building scalable APIs.") -> bytes:
    """Helper to generate in-memory valid PDF with selectable text."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_scanned_style_pdf() -> bytes:
    """Helper to generate in-memory PDF with an embedded image but zero selectable text."""
    doc = pymupdf.open()
    page = doc.new_page()
    # Create a small RGB pixmap and insert as image
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), 1)
    pix.clear_with(200)
    page.insert_image(pymupdf.Rect(50, 50, 250, 250), stream=pix.tobytes())
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_pdf_with_image_and_text(
    text: str = "Jane Doe\nSenior Backend Engineer\nOver 8 years experience building scalable APIs in Python and FastAPI.",
) -> bytes:
    """Helper to generate in-memory PDF with an embedded image and selectable text."""
    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), 1)
    pix.clear_with(200)
    page.insert_image(pymupdf.Rect(50, 50, 100, 100), stream=pix.tobytes())
    page.insert_text((50, 150), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_blank_pdf() -> bytes:
    """Helper to generate a PDF with a page but no text and no images."""
    doc = pymupdf.open()
    doc.new_page()
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(paragraphs: list[str] = None) -> bytes:
    """Helper to generate in-memory valid DOCX with paragraphs."""
    if paragraphs is None:
        paragraphs = [
            "Jane Doe - Principal Software Architect",
            "Extensive experience in distributed microservices, Python, TypeScript, and cloud infrastructure.",
        ]
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_table_docx() -> bytes:
    """Helper to generate in-memory DOCX containing a table."""
    doc = docx.Document()
    doc.add_paragraph("Technical Competencies and Career Experience Summary")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Experience Level"
    table.cell(1, 0).text = "Python & FastAPI"
    table.cell(1, 1).text = "Expert (7+ years)"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_empty_docx() -> bytes:
    """Helper to generate an empty DOCX without any text."""
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# 1. Valid text-based PDF
def test_extract_valid_pdf():
    pdf_bytes = create_sample_pdf()
    response = client.post(
        "/api/resume/extract",
        files={"file": ("candidate_resume.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "candidate_resume.pdf"
    assert data["file_type"] == "pdf"
    assert data["file_size_bytes"] == len(pdf_bytes)
    assert "Jane Doe" in data["extracted_text"]
    assert data["character_count"] == len(data["extracted_text"])


# 2. Valid DOCX with paragraphs
def test_extract_valid_docx_paragraphs():
    docx_bytes = create_sample_docx()
    response = client.post(
        "/api/resume/extract",
        files={"file": ("candidate_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "candidate_resume.docx"
    assert data["file_type"] == "docx"
    assert data["file_size_bytes"] == len(docx_bytes)
    assert "Principal Software Architect" in data["extracted_text"]
    assert data["character_count"] == len(data["extracted_text"])


# 3. DOCX containing a table
def test_extract_docx_with_table():
    docx_bytes = create_table_docx()
    response = client.post(
        "/api/resume/extract",
        files={"file": ("resume_with_table.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert "Python & FastAPI" in data["extracted_text"]
    assert "Expert (7+ years)" in data["extracted_text"]


# 4. Empty upload (0 bytes)
def test_extract_empty_upload():
    response = client.post(
        "/api/resume/extract",
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "EMPTY_FILE"


# 5. Unsupported extension (.txt)
def test_extract_unsupported_extension():
    response = client.post(
        "/api/resume/extract",
        files={"file": ("resume.txt", b"Plain text resume content", "text/plain")},
    )
    assert response.status_code == 415
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "UNSUPPORTED_FILE_TYPE"


# 6. File type / content mismatch
def test_extract_pdf_content_mismatch():
    fake_pdf = b"This is plain text with no PDF header at all."
    response = client.post(
        "/api/resume/extract",
        files={"file": ("spoofed.pdf", fake_pdf, "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "FILE_TYPE_MISMATCH"


def test_extract_docx_content_mismatch():
    fake_docx = b"This is plain text without the PK zip header."
    response = client.post(
        "/api/resume/extract",
        files={"file": ("spoofed.docx", fake_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "FILE_TYPE_MISMATCH"


# 7. Upload exceeding configured size limit
def test_extract_file_too_large():
    # Override settings to a small 200-byte limit to test streaming size guard
    def override_settings():
        return Settings(max_upload_size_bytes=200)

    app.dependency_overrides[get_settings] = override_settings
    try:
        oversized_pdf = create_sample_pdf()
        response = client.post(
            "/api/resume/extract",
            files={"file": ("oversized.pdf", oversized_pdf, "application/pdf")},
        )
        assert response.status_code == 413
        data = response.json()
        assert "error" in data
        assert data["error"]["code"] == "FILE_TOO_LARGE"
    finally:
        app.dependency_overrides.clear()


# 8. Corrupt PDF
def test_extract_corrupt_pdf():
    # Starts with %PDF header to pass magic byte check, but contains corrupt binary data
    corrupt_pdf = b"%PDF-1.4\ncorrupt binary junk garbage data that fails parsing"
    response = client.post(
        "/api/resume/extract",
        files={"file": ("corrupt.pdf", corrupt_pdf, "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "CORRUPT_DOCUMENT"


# 9. Corrupt DOCX
def test_extract_corrupt_docx():
    # Starts with PK\x03\x04 to pass magic byte check, but is not a valid zip archive
    corrupt_docx = b"PK\x03\x04 corrupt zip archive contents"
    response = client.post(
        "/api/resume/extract",
        files={"file": ("corrupt.docx", corrupt_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "CORRUPT_DOCUMENT"


# 10. Empty DOCX without usable text
def test_extract_empty_docx():
    empty_docx = create_empty_docx()
    response = client.post(
        "/api/resume/extract",
        files={"file": ("empty_resume.docx", empty_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "NO_EXTRACTABLE_TEXT"


# 11. PDF with no usable text & scanned/image-only PDF
def test_extract_blank_pdf():
    blank_pdf = create_blank_pdf()
    response = client.post(
        "/api/resume/extract",
        files={"file": ("blank.pdf", blank_pdf, "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "NO_EXTRACTABLE_TEXT"


def test_extract_scanned_image_only_pdf():
    scanned_pdf = create_scanned_style_pdf()
    response = client.post(
        "/api/resume/extract",
        files={"file": ("scanned_resume.pdf", scanned_pdf, "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == "SCANNED_PDF_DETECTED"
    assert "text-selectable PDF" in data["error"]["message"]


# 12. Successful response schema and character count accuracy
def test_extract_response_schema_and_character_count():
    text = "Full Stack Engineer with extensive experience in Next.js, React, and Python FastAPI backend development."
    pdf_bytes = create_sample_pdf(text=text)
    response = client.post(
        "/api/resume/extract",
        files={"file": ("my_resume.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert set(data.keys()) == {"filename", "file_type", "file_size_bytes", "extracted_text", "character_count"}
    assert data["character_count"] == len(data["extracted_text"])
    assert data["character_count"] > 0


# 13. Confirm all failure scenarios strictly use standard error envelope
@pytest.mark.parametrize(
    "filename, content, content_type, expected_status, expected_code",
    [
        ("empty.pdf", b"", "application/pdf", 400, "EMPTY_FILE"),
        ("invalid.xyz", b"abc", "application/octet-stream", 415, "UNSUPPORTED_FILE_TYPE"),
        ("fake.pdf", b"Not a real PDF", "application/pdf", 400, "FILE_TYPE_MISMATCH"),
    ],
)
def test_standard_error_envelope_structure(filename, content, content_type, expected_status, expected_code):
    response = client.post(
        "/api/resume/extract",
        files={"file": (filename, content, content_type)},
    )
    assert response.status_code == expected_status
    data = response.json()
    assert "error" in data
    assert isinstance(data["error"]["code"], str)
    assert data["error"]["code"] == expected_code
    assert isinstance(data["error"]["message"], str)


# 14. Regression: Application-exported PDF extract & tailor roundtrip
def test_extract_and_tailor_application_generated_pdf_roundtrip():
    """Verify that a PDF exported by our application can be extracted and then tailored."""
    resume_data = ResumeData.model_validate(MOCK_TAILORED_PAYLOAD["tailored_resume"])
    exported_pdf_bytes = generate_resume_pdf(resume_data)

    # Step 1: Extract text from application-generated PDF
    response = client.post(
        "/api/resume/extract",
        files={"file": ("Jane_Doe_Resume.pdf", exported_pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200, f"Extract failed: {response.text}"
    extract_data = response.json()
    assert extract_data["filename"] == "Jane_Doe_Resume.pdf"
    assert extract_data["file_type"] == "pdf"
    assert extract_data["character_count"] > 300
    assert "Jane Doe" in extract_data["extracted_text"]
    assert "Acme Systems" in extract_data["extracted_text"]
    assert "San Francisco" in extract_data["extracted_text"]

    # Step 2: Feed extracted text into tailoring endpoint with mocked Gemini client
    app.dependency_overrides[get_settings] = get_mock_settings
    try:
        mock_client = MockGenAIClient(response_text=json.dumps(MOCK_TAILORED_PAYLOAD))
        with patch("app.services.llm.client.genai.Client", return_value=mock_client):
            tailor_resp = client.post(
                "/api/resume/tailor",
                json={
                    "extracted_text": extract_data["extracted_text"],
                    "job_description": "We are seeking a Senior Backend Engineer experienced in Python, FastAPI, and Docker.",
                },
            )
            assert tailor_resp.status_code == 200, f"Tailor failed: {tailor_resp.text}"
            tailor_data = tailor_resp.json()
            assert "tailored_resume" in tailor_data
            assert tailor_data["alignment_score"] == 88
    finally:
        app.dependency_overrides.clear()


# 15. Regression: PDF with embedded image and selectable text is NOT rejected as scanned
def test_extract_pdf_with_image_and_selectable_text_succeeds():
    """Verify that a resume with a profile image/logo and selectable text is accepted."""
    pdf_bytes = create_pdf_with_image_and_text()
    response = client.post(
        "/api/resume/extract",
        files={"file": ("resume_with_photo.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["character_count"] >= 30
    assert "Jane Doe" in data["extracted_text"]


# 16. Regression: PDF with insufficient text (< 30 chars) and no images returns NO_EXTRACTABLE_TEXT
def test_extract_pdf_insufficient_text_without_images_returns_no_extractable_text():
    """Verify that a text-based PDF with too few characters is rejected with NO_EXTRACTABLE_TEXT."""
    pdf_bytes = create_sample_pdf(text="Short text")
    response = client.post(
        "/api/resume/extract",
        files={"file": ("short.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert data["error"]["code"] == "NO_EXTRACTABLE_TEXT"
    assert "sufficient readable text" in data["error"]["message"]

