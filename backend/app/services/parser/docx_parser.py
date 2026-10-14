import io
import docx
from app.core.exceptions import AppException


def extract_text_from_docx(file_bytes: bytes, min_chars: int = 30) -> str:
    """Extract readable text from DOCX bytes across paragraphs and tables.

    Preserves paragraph breaks and structured cell content.
    Handles corrupt, empty, and textless documents with clear errors.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception:
        raise AppException(
            status_code=400,
            code="CORRUPT_DOCUMENT",
            message="The uploaded DOCX document is corrupt or could not be read.",
        )

    content_blocks = []

    # 1. Extract text from paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content_blocks.append(text)

    # 2. Extract text from tables, preserving structured cell content
    for table in doc.tables:
        for row in table.rows:
            row_cells_text = []
            seen_cells = set()
            for cell in row.cells:
                # Deduplicate identical merged cells within the same row
                if cell not in seen_cells:
                    seen_cells.add(cell)
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_cells_text.append(cell_text)
            if row_cells_text:
                content_blocks.append(" | ".join(row_cells_text))

    extracted_text = "\n\n".join(content_blocks).strip()
    non_whitespace_count = len("".join(extracted_text.split()))

    if non_whitespace_count < min_chars:
        raise AppException(
            status_code=400,
            code="NO_EXTRACTABLE_TEXT",
            message="The uploaded DOCX document does not contain sufficient readable text.",
            details={
                "selectable_chars": non_whitespace_count,
                "min_required_chars": min_chars,
            },
        )

    return extracted_text
