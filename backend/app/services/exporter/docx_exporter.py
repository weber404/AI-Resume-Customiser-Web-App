import io
from typing import Union

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas.resume import ResumeData, TailoredResumeResponse

# Palette definitions
COLOR_PRIMARY = RGBColor(15, 23, 42)    # Slate 900
COLOR_SUBTLE = RGBColor(71, 85, 105)    # Slate 600
COLOR_ACCENT = RGBColor(3, 105, 161)    # Sky 700


def _set_paragraph_bottom_border(paragraph):
    """Add a subtle bottom border to section heading paragraph for ATS readability."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")  # 0.75 pt
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "0284C7")  # Sky 600
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_section_heading(doc: docx.Document, title: str):
    """Add standardized uppercase ATS section header with keep_with_next set."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    _set_paragraph_bottom_border(p)

    run = p.add_run(title.upper())
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_ACCENT
    return p


def generate_resume_docx(data: Union[TailoredResumeResponse, ResumeData]) -> bytes:
    """Build a professional, single-column ATS-friendly DOCX resume in memory."""
    if isinstance(data, TailoredResumeResponse):
        resume = data.tailored_resume
    else:
        resume = data

    doc = docx.Document()

    # Configure 0.5-inch standard margins
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    # Set base Normal style font
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = COLOR_PRIMARY

    # 1. Contact Info Header
    if resume.contact:
        c = resume.contact
        if c.full_name:
            p_name = doc.add_paragraph()
            p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_name.paragraph_format.space_before = Pt(0)
            p_name.paragraph_format.space_after = Pt(2)
            p_name.paragraph_format.keep_with_next = True
            r_name = p_name.add_run(c.full_name)
            r_name.font.name = "Arial"
            r_name.font.size = Pt(20)
            r_name.font.bold = True
            r_name.font.color.rgb = COLOR_PRIMARY

        contact_parts = []
        if c.location:
            contact_parts.append(c.location)
        if c.email:
            contact_parts.append(c.email)
        if c.phone:
            contact_parts.append(c.phone)
        if c.linkedin_url:
            contact_parts.append(c.linkedin_url)
        if c.github_url:
            contact_parts.append(c.github_url)
        if c.portfolio_url:
            contact_parts.append(c.portfolio_url)

        if contact_parts:
            p_contact = doc.add_paragraph()
            p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_contact.paragraph_format.space_before = Pt(0)
            p_contact.paragraph_format.space_after = Pt(8)
            p_contact.paragraph_format.keep_with_next = True
            r_contact = p_contact.add_run(" • ".join(contact_parts))
            r_contact.font.name = "Arial"
            r_contact.font.size = Pt(8.5)
            r_contact.font.color.rgb = COLOR_SUBTLE

    # 2. Professional Summary
    if resume.professional_summary and resume.professional_summary.strip():
        _add_section_heading(doc, "Professional Summary")
        p_summary = doc.add_paragraph()
        p_summary.paragraph_format.space_before = Pt(2)
        p_summary.paragraph_format.space_after = Pt(4)
        p_summary.paragraph_format.line_spacing = Pt(13)
        r_sum = p_summary.add_run(resume.professional_summary.strip())
        r_sum.font.size = Pt(9.5)

    # 3. Skills
    if resume.skills and any(cat.skills for cat in resume.skills):
        _add_section_heading(doc, "Skills")
        for cat in resume.skills:
            if cat.skills:
                p_skill = doc.add_paragraph()
                p_skill.paragraph_format.space_before = Pt(1)
                p_skill.paragraph_format.space_after = Pt(2)
                p_skill.paragraph_format.line_spacing = Pt(13)
                r_cat = p_skill.add_run(f"{cat.category}: ")
                r_cat.font.bold = True
                r_cat.font.size = Pt(9.5)
                r_slist = p_skill.add_run(", ".join(cat.skills))
                r_slist.font.size = Pt(9.5)

    # 4. Work Experience
    if resume.experience:
        _add_section_heading(doc, "Work Experience")
        for exp in resume.experience:
            # Title & Company
            p_role = doc.add_paragraph()
            p_role.paragraph_format.space_before = Pt(4)
            p_role.paragraph_format.space_after = Pt(1)
            p_role.paragraph_format.keep_with_next = True

            r_title = p_role.add_run(exp.title)
            r_title.font.bold = True
            r_title.font.size = Pt(10)
            r_co = p_role.add_run(f" — {exp.company}")
            r_co.font.size = Pt(10)

            # Dates & Location
            date_range = exp.start_date
            if exp.is_current or not exp.end_date:
                date_range += " – Present"
            elif exp.end_date:
                date_range += f" – {exp.end_date}"

            sub_parts = [date_range]
            if exp.location:
                sub_parts.append(exp.location)

            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.keep_with_next = True
            r_sub = p_sub.add_run(" | ".join(sub_parts))
            r_sub.font.size = Pt(8.5)
            r_sub.font.italic = True
            r_sub.font.color.rgb = COLOR_SUBTLE

            # Bullet points
            for bullet in exp.bullet_points:
                if bullet.strip():
                    p_bullet = doc.add_paragraph(style="List Bullet")
                    p_bullet.paragraph_format.space_before = Pt(1)
                    p_bullet.paragraph_format.space_after = Pt(2)
                    p_bullet.paragraph_format.line_spacing = Pt(12)
                    r_b = p_bullet.add_run(bullet.strip())
                    r_b.font.size = Pt(9)

            # Key Technologies
            if exp.technologies_used:
                p_tech = doc.add_paragraph()
                p_tech.paragraph_format.left_indent = Inches(0.2)
                p_tech.paragraph_format.space_before = Pt(1)
                p_tech.paragraph_format.space_after = Pt(4)
                r_t_lbl = p_tech.add_run("Key Technologies: ")
                r_t_lbl.font.size = Pt(8.5)
                r_t_lbl.font.italic = True
                r_t_lbl.font.color.rgb = COLOR_SUBTLE
                r_t_val = p_tech.add_run(", ".join(exp.technologies_used))
                r_t_val.font.size = Pt(8.5)
                r_t_val.font.color.rgb = COLOR_SUBTLE

    # 5. Projects
    if resume.projects:
        _add_section_heading(doc, "Projects")
        for proj in resume.projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = Pt(4)
            p_proj.paragraph_format.space_after = Pt(1)
            p_proj.paragraph_format.keep_with_next = True

            r_p_title = p_proj.add_run(proj.title)
            r_p_title.font.bold = True
            r_p_title.font.size = Pt(10)
            if proj.url:
                r_p_url = p_proj.add_run(f" • {proj.url}")
                r_p_url.font.size = Pt(8.5)
                r_p_url.font.italic = True
                r_p_url.font.color.rgb = COLOR_SUBTLE

            if proj.description and proj.description.strip():
                p_desc = doc.add_paragraph()
                p_desc.paragraph_format.space_before = Pt(0)
                p_desc.paragraph_format.space_after = Pt(2)
                r_d = p_desc.add_run(proj.description.strip())
                r_d.font.size = Pt(9)

            for bullet in proj.bullet_points:
                if bullet.strip():
                    p_b = doc.add_paragraph(style="List Bullet")
                    p_b.paragraph_format.space_before = Pt(1)
                    p_b.paragraph_format.space_after = Pt(2)
                    p_b.paragraph_format.line_spacing = Pt(12)
                    r_pb = p_b.add_run(bullet.strip())
                    r_pb.font.size = Pt(9)

            if proj.technologies_used:
                p_ptech = doc.add_paragraph()
                p_ptech.paragraph_format.left_indent = Inches(0.2)
                p_ptech.paragraph_format.space_before = Pt(1)
                p_ptech.paragraph_format.space_after = Pt(4)
                r_pt_lbl = p_ptech.add_run("Technologies: ")
                r_pt_lbl.font.size = Pt(8.5)
                r_pt_lbl.font.italic = True
                r_pt_lbl.font.color.rgb = COLOR_SUBTLE
                r_pt_val = p_ptech.add_run(", ".join(proj.technologies_used))
                r_pt_val.font.size = Pt(8.5)
                r_pt_val.font.color.rgb = COLOR_SUBTLE

    # 6. Education
    if resume.education:
        _add_section_heading(doc, "Education")
        for edu in resume.education:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = Pt(4)
            p_edu.paragraph_format.space_after = Pt(1)
            p_edu.paragraph_format.keep_with_next = True

            deg_text = edu.degree
            if edu.field_of_study:
                deg_text += f" in {edu.field_of_study}"
            r_deg = p_edu.add_run(deg_text)
            r_deg.font.bold = True
            r_deg.font.size = Pt(10)
            r_inst = p_edu.add_run(f" — {edu.institution}")
            r_inst.font.size = Pt(10)

            sub_parts = []
            if edu.start_date or edu.end_date:
                d_str = ""
                if edu.start_date:
                    d_str += f"{edu.start_date} – "
                d_str += edu.end_date or "Present"
                sub_parts.append(d_str)
            if edu.gpa:
                sub_parts.append(f"GPA: {edu.gpa}")

            if sub_parts:
                p_esub = doc.add_paragraph()
                p_esub.paragraph_format.space_before = Pt(0)
                p_esub.paragraph_format.space_after = Pt(2)
                p_esub.paragraph_format.keep_with_next = True
                r_es = p_esub.add_run(" | ".join(sub_parts))
                r_es.font.size = Pt(8.5)
                r_es.font.color.rgb = COLOR_SUBTLE

            for highlight in edu.highlights:
                if highlight.strip():
                    p_hb = doc.add_paragraph(style="List Bullet")
                    p_hb.paragraph_format.space_before = Pt(1)
                    p_hb.paragraph_format.space_after = Pt(2)
                    r_h = p_hb.add_run(highlight.strip())
                    r_h.font.size = Pt(9)

    # 7. Certifications
    if resume.certifications:
        _add_section_heading(doc, "Certifications")
        for cert in resume.certifications:
            p_cert = doc.add_paragraph()
            p_cert.paragraph_format.space_before = Pt(3)
            p_cert.paragraph_format.space_after = Pt(1)
            p_cert.paragraph_format.keep_with_next = True

            r_cname = p_cert.add_run(cert.name)
            r_cname.font.bold = True
            r_cname.font.size = Pt(9.5)
            r_ciss = p_cert.add_run(f" — {cert.issuer}")
            r_ciss.font.size = Pt(9.5)
            if cert.issue_date:
                r_cdate = p_cert.add_run(f" ({cert.issue_date})")
                r_cdate.font.size = Pt(8.5)
                r_cdate.font.color.rgb = COLOR_SUBTLE

            if cert.credential_id or cert.url:
                p_cdet = doc.add_paragraph()
                p_cdet.paragraph_format.space_before = Pt(0)
                p_cdet.paragraph_format.space_after = Pt(2)
                details = []
                if cert.credential_id:
                    details.append(f"Credential ID: {cert.credential_id}")
                if cert.url:
                    details.append(cert.url)
                r_cd = p_cdet.add_run(" | ".join(details))
                r_cd.font.size = Pt(8)
                r_cd.font.color.rgb = COLOR_SUBTLE

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
